import asyncio
import copy
import time

import httpx
import pytest

from autoreview import settings, store
from autoreview.pipeline import Runner
from autoreview.reasons import load_reasons

FILLER = "فروشگاه آنلاین پوشاک با ارسال سریع به سراسر کشور و ضمانت بازگشت کالا. " * 8


def home_html(host):
    return (f"<html><body><h1>{host}</h1><p>{FILLER}</p><footer>تلفن تماس ۰۲۱۱۲۳۴۵۶۷۸ "
            f"<a href='https://trustseal.enamad.ir/?id={host}&code=c'>enamad</a></footer></body></html>")


def profile_html(host):
    return (f"<div>{host}</div><div>صاحب امتیاز :</div><div>علی رضایی</div><div>تاریخ اعتبار :</div><div>معتبر تا تاریخ 1410/01/01</div>"
            "<table><tbody><tr><td>1</td><td>فروش پوشاک، کیف، کفش و محصولات چرمی</td><td></td><td>-</td><td>-</td><td>-</td><td>تایید شده</td></tr></tbody></table>")


PRODUCT_SITEMAP = "<urlset>" + "".join(f"<url><loc>https://HOST/product/p{i}</loc></url>" for i in range(80)) + "</urlset>"
SHOPS = {"shop.ir", "shop2.ir", "shop3.ir"}


def make_handler(delay=0.0):
    async def handler(request: httpx.Request):
        if delay:
            await asyncio.sleep(delay)
        h, path = request.url.host, request.url.path
        if h == "enamad.ir":
            if request.method == "POST":
                dom = request.content.decode().split("=", 1)[1]
                return httpx.Response(200, json={"id": dom, "code": "c", "enamad_status": 1, "expdate": "1410/01/01"})
            return httpx.Response(200, text="home", headers={"set-cookie": "a=b"})
        if h == "trustseal.enamad.ir":
            return httpx.Response(200, text=profile_html(request.url.params["id"]))
        if h in SHOPS:
            if path == "/robots.txt":
                return httpx.Response(200, text=f"Sitemap: https://{h}/product-sitemap.xml")
            if path == "/product-sitemap.xml":
                return httpx.Response(200, text=PRODUCT_SITEMAP.replace("HOST", h))
            if path.startswith("/product/"):
                return httpx.Response(200, text="<form><button name='add-to-cart'>افزودن به سبد خرید</button></form>")
            if path == "/":
                return httpx.Response(200, text=home_html(h))
            return httpx.Response(404)
        raise httpx.ConnectError("[Errno 11001] getaddrinfo failed")
    return handler


def factory(delay=0.0):
    return lambda: httpx.AsyncClient(transport=httpx.MockTransport(make_handler(delay)))


def row(smr, site, holder="علی رضایی"):
    return {"smr": smr, "site": site, "category": "مد و پوشاک", "created_at": "1405/07/08", "account_holder": holder,
            "owner_name": "علی", "owner_family": "رضایی", "has_online": "true", "has_instore": "false"}


NBO_OK = [{"id": "OLD", "site": "dup.ir"}]
CRM_OK = [{"id": "MRG-9", "site": "crm-only.ir"}]


def results(home, run_id):
    db = store.connect(home / "t.db")
    return {x["smr"]: x for x in store.results_of(db, run_id)}


def test_batch_end_to_end_offline(isolated_profile):
    r = Runner(isolated_profile / "t.db", reasons=copy.deepcopy(load_reasons()), client_factory=factory())
    rows = [row("A", "shop.ir"), row("B", "dead.ir"), row("C", "shop2.ir", holder="فرد دیگر"), row("D", "dup.ir"),
            row("E", "https://instagram.com/my_shop")]
    r.start(rows, approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t1")
    r.join(60)
    assert r.progress.state == "finished" and r.progress.done == 5
    got = results(isolated_profile, "t1")
    assert got["A"]["action"] == "APPROVE", got["A"]["notes"]
    assert got["B"]["action"] == "EDIT" and got["B"]["reason_codes"] == ["INVALID_URL"]   # Action Test 4: a dead address => EDIT
    # Action Test 4 checks the registrant vs the bank account before the enamad owner
    assert got["C"]["action"] == "EDIT" and got["C"]["reason_codes"] == ["REGISTRANT_NAME_AND_BANK_ACCOUNT_OWNER_MISMATCH"]
    assert got["D"]["action"] == "CANCEL" and got["D"]["reason_codes"] == ["DUPLICATE_REQUEST"] and "NBO: OLD" in got["D"]["notes"][0]
    assert got["E"]["action"] == "MANUAL" and "shared platform" in got["E"]["notes"][0]
    assert all(x["duration_ms"] is not None for x in got.values())
    runs = store.list_runs(store.connect(isolated_profile / "t.db"))
    assert runs[0]["run_id"] == "t1" and runs[0]["state"] == "finished" and runs[0]["counts"]["APPROVE"] == 1


def test_two_pending_requests_for_one_site_go_to_a_person(isolated_profile):
    r = Runner(isolated_profile / "t.db", client_factory=factory())
    batch = [row("A", "shop.ir")]
    r.start(batch, approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t5", pending_all=batch + [row("Z", "https://www.shop.ir/")])
    r.join(60)
    got = results(isolated_profile, "t5")
    assert got["A"]["action"] == "MANUAL" and "Z" in got["A"]["notes"][0]


def test_a_run_needs_both_nbo_and_crm_and_a_crm_only_duplicate_is_caught(isolated_profile):
    r = Runner(isolated_profile / "t.db", client_factory=factory())
    with pytest.raises(ValueError):
        r.start([row("A", "shop.ir")], approved_nbo=NBO_OK, approved_crm=[], run_id="t3")
    with pytest.raises(ValueError):
        r.start([row("A", "shop.ir")], approved_nbo=[], approved_crm=CRM_OK, run_id="t3")
    r.start([row("A", "shop.ir")], approved_nbo=NBO_OK, approved_crm=[{"id": "MRG-5", "site": "https://www.shop.ir/"}], run_id="t4")
    r.join(60)
    db = store.connect(isolated_profile / "t.db")
    got = store.results_of(db, "t4")[0]
    assert got["action"] == "CANCEL" and "CRM: MRG-5" in got["notes"][0]
    assert store.run_sources(db, "t4") == {"run": "t4", "nbo_approved": 1, "crm_approved": 1}


def test_stop_ends_the_run_cleanly(isolated_profile):
    r = Runner(isolated_profile / "t.db", client_factory=factory(delay=0.02))
    r.start([row(str(i), "shop.ir") for i in range(50)], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t2", pending_all=[])
    time.sleep(0.3)
    r.stop()
    r.join(60)
    assert r.progress.state == "stopped" and r.progress.done < 50
    assert store.list_runs(store.connect(isolated_profile / "t.db"))[0]["state"] == "stopped"


def test_pause_really_stops_starting_new_requests(isolated_profile):
    rules = settings.load_rules()
    rules["runtime"]["concurrency"] = 2
    r = Runner(isolated_profile / "t.db", rules=rules, client_factory=factory(delay=0.03))
    r.start([row(str(i), "shop.ir") for i in range(20)], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t6", pending_all=[])
    time.sleep(0.4)
    r.pause()
    time.sleep(2.0)                                   # requests already in flight finish
    frozen = r.progress.done
    time.sleep(1.0)
    assert r.progress.done == frozen and r.progress.state == "paused"
    r.resume()
    r.join(120)
    assert r.progress.state == "finished" and r.progress.done == 20


def test_a_request_that_takes_too_long_goes_to_a_person(isolated_profile):
    rules = settings.load_rules()
    rules["runtime"]["request_deadline_seconds"] = 0.3
    r = Runner(isolated_profile / "t.db", rules=rules, client_factory=factory(delay=0.2))
    r.start([row("A", "shop.ir")], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t7")
    r.join(60)
    got = results(isolated_profile, "t7")
    assert got["A"]["action"] == "MANUAL" and got["A"]["notes"][0].startswith("timed out")


def test_a_seal_drawn_by_javascript_is_seen_by_the_hidden_browser(isolated_profile):
    base = make_handler()

    async def handler(request):                       # shop3.ir: the seal is NOT in the downloaded HTML (JavaScript draws it)
        if request.url.host == "shop3.ir" and request.url.path == "/":
            return httpx.Response(200, text=home_html("shop3.ir").replace("trustseal.enamad.ir", "cdn.example.ir"))
        return await base(request)
    rendered = []

    async def render(url):
        rendered.append(url)
        return home_html("shop3.ir")                  # after JavaScript ran, the seal is on the page

    def mk():
        return httpx.AsyncClient(transport=httpx.MockTransport(handler))
    rules = settings.load_rules()
    rules["checks"]["enamad_on_site"]["enabled"] = True     # not an Action Test 4 rule: off by default, switched on here
    no_browser = Runner(isolated_profile / "t.db", rules=rules, client_factory=mk)
    no_browser.start([row("A", "shop3.ir")], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="r0")
    no_browser.join(60)
    assert results(isolated_profile, "r0")["A"]["action"] == "MANUAL"
    with_browser = Runner(isolated_profile / "t.db", rules=rules, client_factory=mk, render=render)
    with_browser.start([row("A", "shop3.ir")], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="r1")
    with_browser.join(60)
    got = results(isolated_profile, "r1")["A"]
    assert got["action"] == "APPROVE" and rendered == ["https://shop3.ir"]


def test_the_hidden_browser_is_not_used_when_it_cannot_change_the_outcome(isolated_profile):
    calls = []

    async def render(url):
        calls.append(url)
        return None
    r = Runner(isolated_profile / "t.db", client_factory=factory(), render=render)
    r.start([row("A", "shop.ir"), row("D", "dup.ir"), row("C", "shop2.ir", holder="فرد دیگر")], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="r2")
    r.join(60)
    assert calls == []


def test_dashboard_numbers(isolated_profile):
    rules = settings.load_rules()
    rules["unreachable_site_action"]["value"] = "MANUAL"     # keeps one 'why manual' cause in the numbers
    r = Runner(isolated_profile / "t.db", rules=rules, client_factory=factory())
    r.start([row("A", "shop.ir"), row("D", "dup.ir"), row("B", "dead.ir")], approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t8")
    r.join(60)
    db = store.connect(isolated_profile / "t.db")
    t = store.totals(db, run_id="t8")
    assert t["total"] == 3 and t["counts"]["CANCEL"] == 1 and t["avg_seconds"] is not None
    assert store.reason_counts(db, run_id="t8") == [("DUPLICATE_REQUEST", 1)]
    assert store.manual_causes(db, run_id="t8")[0][1] == 1
    days = store.daily_counts(db, days=7)
    assert len(days) == 7 and sum(days[-1][1].values()) == 3


def test_export_has_three_tabs_and_never_writes_formulas(tmp_path):
    from openpyxl import load_workbook
    from autoreview.export import write_xlsx
    res = [dict(smr="A", site="=HYPERLINK(\"http://x\")", category="c", created_at="d", action="MANUAL", reason_keys=[], reason_codes=[],
                notes=["could not determine: add to cart"], decided_at="", duration_ms=1200)]
    write_xlsx(tmp_path / "o.xlsx", res)
    wb = load_workbook(tmp_path / "o.xlsx")
    assert wb.sheetnames[:3] == ["نتایج", "صف دستی", "خلاصه"] and wb["صف دستی"].max_row == 2
    cell = wb["نتایج"]["B2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")              # stored as text, not as a formula


def test_https_as_the_old_code_the_address_registered_in_nbo_decides(isolated_profile):
    r = Runner(isolated_profile / "t.db", reasons=copy.deepcopy(load_reasons()), client_factory=factory())
    rows = [row("A", "http://shop.ir"), row("B", "shop2.ir"), row("C", "https://shop3.ir")]
    r.start(rows, approved_nbo=NBO_OK, approved_crm=CRM_OK, run_id="t9")
    r.join(60)
    got = results(isolated_profile, "t9")
    assert got["A"]["action"] == "EDIT" and got["A"]["reason_codes"] == ["INVALID_URL"]     # typed http:// = 'URL wrong'
    assert got["B"]["action"] == "APPROVE", got["B"]["notes"]                               # no scheme = https (NBO links it so)
    assert got["C"]["action"] == "APPROVE", got["C"]["notes"]
